"""End-to-end: a PDF on disk becomes a parsed, described paper row.

Exercises the real registrar, the real queue, the real worker handlers and the
real tier-0 parser. Only the GPU tiers are stubbed, because a test suite must
not depend on Marker being installed or Ollama being up.
"""

from __future__ import annotations

import json
import shutil

import pytest

from app.core.config import Settings
from app.models import (
    Job,
    JobKind,
    JobState,
    MarkdownDoc,
    Paper,
    PaperMeta,
    PaperStatus,
)
from app.services.ingest.registrar import Registration, register_document
from app.workers.handlers import handle_metadata, handle_parse
from app.workers.queue import claim_next, complete


@pytest.fixture
def settings(tmp_path) -> Settings:
    return Settings(
        data_dir=tmp_path,
        library_dir=tmp_path / "library",
        markdown_dir=tmp_path / "markdown",
        vectors_dir=tmp_path / "vectors",
        models_dir=tmp_path / "models",
        db_path=tmp_path / "t.db",
        enrichment_enabled=False,
    )


@pytest.fixture
def library(settings, pdf_fixtures):
    settings.ensure_dirs()
    placed = {}
    for name in ("with_identifiers.pdf", "clean_single_column.pdf"):
        dest = settings.library_dir / name
        shutil.copy(pdf_fixtures[name], dest)
        placed[name] = dest
    return placed


def run_stage(sf, kind: JobKind, settings: Settings, handler) -> int:
    """Drain one stage the way the worker does."""
    done = 0
    while True:
        with sf() as session:
            job = claim_next(session, kinds=[kind])
            if job is None:
                return done
            handler(session, job, settings)
            complete(session, job)
            session.commit()
            done += 1


def test_pdf_becomes_a_parsed_paper(sf, settings, library):
    with sf() as s:
        result = register_document(s, library["with_identifiers.pdf"])
        s.commit()
        assert result.outcome is Registration.CREATED
        paper_id = result.paper.id

    assert run_stage(sf, JobKind.PARSE, settings, handle_parse) == 1

    with sf() as s:
        paper = s.get(Paper, paper_id)
        md = s.get(MarkdownDoc, paper_id)
        assert paper.status == PaperStatus.PARSED
        assert md is not None
        assert md.tier == 0, "a clean text layer must not reach the GPU"
        assert md.char_count > 500
        from pathlib import Path

        assert Path(md.md_path).exists()


def test_metadata_stage_extracts_identifiers(sf, settings, library):
    with sf() as s:
        paper_id = register_document(s, library["with_identifiers.pdf"]).paper.id
        s.commit()

    run_stage(sf, JobKind.PARSE, settings, handle_parse)
    assert run_stage(sf, JobKind.METADATA, settings, handle_metadata) == 1

    with sf() as s:
        meta = s.get(PaperMeta, paper_id)
        paper = s.get(Paper, paper_id)
        assert meta.doi == "10.1145/3292500.3330701"
        assert meta.arxiv_id == "2401.01234v2"
        # Metadata no longer ends the pipeline: embedding follows it.
        assert paper.status == PaperStatus.PARSED
        # The work key must have been upgraded from the content hash.
        assert paper.work_key == "doi:10.1145/3292500.3330701"


def test_metadata_hands_off_to_embedding(sf, settings, library):
    """The stage boundary that carries a paper into M2."""
    with sf() as s:
        paper_id = register_document(s, library["with_identifiers.pdf"]).paper.id
        s.commit()
    run_stage(sf, JobKind.PARSE, settings, handle_parse)
    run_stage(sf, JobKind.METADATA, settings, handle_metadata)

    with sf() as s:
        queued = (
            s.query(Job)
            .filter(Job.kind == JobKind.EMBED, Job.paper_id == paper_id)
            .count()
        )
        assert queued == 1


def test_parse_enqueues_metadata(sf, settings, library):
    with sf() as s:
        register_document(s, library["with_identifiers.pdf"])
        s.commit()
    run_stage(sf, JobKind.PARSE, settings, handle_parse)

    with sf() as s:
        queued = (
            s.query(Job)
            .filter(Job.kind == JobKind.METADATA, Job.state == JobState.QUEUED)
            .count()
        )
        assert queued == 1


def test_reingesting_the_same_bytes_is_a_noop(sf, settings, library):
    with sf() as s:
        first = register_document(s, library["with_identifiers.pdf"])
        s.commit()
    with sf() as s:
        second = register_document(s, library["with_identifiers.pdf"])
        s.commit()

    assert second.outcome is Registration.DUPLICATE_CONTENT
    assert second.existing.id == first.paper.id
    with sf() as s:
        assert s.query(Paper).count() == 1


def test_a_moved_file_follows_rather_than_duplicating(sf, settings, library):
    with sf() as s:
        paper_id = register_document(s, library["with_identifiers.pdf"]).paper.id
        s.commit()

    moved = settings.library_dir / "renamed.pdf"
    library["with_identifiers.pdf"].rename(moved)

    with sf() as s:
        result = register_document(s, moved)
        s.commit()
        assert result.outcome is Registration.MOVED
        assert s.query(Paper).count() == 1
        assert s.get(Paper, paper_id).pdf_path == str(moved.resolve())


def test_quality_report_is_persisted_for_diagnosis(sf, settings, library):
    with sf() as s:
        paper_id = register_document(s, library["clean_single_column.pdf"]).paper.id
        s.commit()
    run_stage(sf, JobKind.PARSE, settings, handle_parse)

    with sf() as s:
        report = json.loads(s.get(MarkdownDoc, paper_id).quality_report_json)
        assert report["passed"] is True
        assert report["degraded"] is False
        assert "chars_per_page" in report["metrics"]


def test_a_missing_pdf_fails_the_job_not_the_worker(sf, settings, library):
    with sf() as s:
        register_document(s, library["clean_single_column.pdf"])
        s.commit()
    library["clean_single_column.pdf"].unlink()

    with sf() as s:
        job = claim_next(s, kinds=[JobKind.PARSE])
        with pytest.raises(FileNotFoundError):
            handle_parse(s, job, settings)


# --- recursive, multi-format ingestion -------------------------------------
#
# A library is a folder tree, not a flat directory, and not everything in it is
# a PDF. These assert the two halves of that: discovery reaches every depth,
# and a non-PDF comes out the far end of the pipeline indistinguishable from a
# PDF.


def test_discovery_reaches_every_subdirectory(settings, pdf_fixtures):
    from app.services.ingest.registrar import library_documents

    settings.ensure_dirs()
    root = settings.library_dir
    (root / "2024" / "neuro").mkdir(parents=True)
    (root / "unsorted").mkdir()

    shutil.copy(pdf_fixtures["with_identifiers.pdf"], root / "top.pdf")
    shutil.copy(pdf_fixtures["clean_single_column.pdf"], root / "2024" / "mid.pdf")
    (root / "2024" / "neuro" / "deep.md").write_text(
        "# Deep\n\nProse.", encoding="utf-8"
    )
    (root / "unsorted" / "notes.txt").write_text("Some notes here.", encoding="utf-8")
    # Not a document, and must not be picked up.
    (root / "unsorted" / "cover.png").write_bytes(b"\x89PNG\r\n\x1a\n")

    found = {p.name for p in library_documents(root)}
    assert found == {"top.pdf", "mid.pdf", "deep.md", "notes.txt"}


def test_a_markdown_file_becomes_a_paper(sf, settings):
    """The whole point of multi-format support, checked end to end."""
    settings.ensure_dirs()
    nested = settings.library_dir / "reading" / "2026"
    nested.mkdir(parents=True)
    source = nested / "review.md"
    source.write_text(
        "# Structural Constraints on Neural Development\n\n"
        "## Abstract\n\n"
        "This review examines how mechanical constraints shape the developing "
        "cortex, drawing together evidence from imaging and modelling work "
        "across several model organisms and developmental stages.\n",
        encoding="utf-8",
    )

    with sf() as s:
        result = register_document(s, source)
        s.commit()
        assert result.outcome is Registration.CREATED
        paper_id = result.paper.id

    assert run_stage(sf, JobKind.PARSE, settings, handle_parse) == 1
    assert run_stage(sf, JobKind.METADATA, settings, handle_metadata) == 1

    with sf() as s:
        paper = s.get(Paper, paper_id)
        md = s.get(MarkdownDoc, paper_id)
        meta = s.get(PaperMeta, paper_id)

        assert paper.status == PaperStatus.PARSED
        assert md.parser == "markdown-passthrough"
        assert md.tier == 0
        # The title must survive to the map, exactly as it would from a PDF.
        assert meta.title == "Structural Constraints on Neural Development"
        assert meta.abstract and "mechanical constraints" in meta.abstract


def test_a_docx_becomes_a_paper(sf, settings):
    pytest.importorskip("docx")
    import docx

    settings.ensure_dirs()
    source = settings.library_dir / "drafts" / "paper.docx"
    source.parent.mkdir(parents=True)
    document = docx.Document()
    document.add_heading("Epigenetic Drift in Long-Lived Cells", level=1)
    document.add_paragraph(
        "We tracked methylation across three decades of samples and found that "
        "drift accumulates in a manner consistent with replicative age rather "
        "than chronological age."
    )
    document.save(str(source))

    with sf() as s:
        paper_id = register_document(s, source).paper.id
        s.commit()

    assert run_stage(sf, JobKind.PARSE, settings, handle_parse) == 1
    run_stage(sf, JobKind.METADATA, settings, handle_metadata)

    with sf() as s:
        md = s.get(MarkdownDoc, paper_id)
        meta = s.get(PaperMeta, paper_id)
        assert md.parser == "python-docx"
        assert meta.title == "Epigenetic Drift in Long-Lived Cells"


def test_the_next_stage_is_queued_for_a_non_pdf(sf, settings):
    """A .txt must not stall: parse has to hand off to metadata like a PDF."""
    settings.ensure_dirs()
    source = settings.library_dir / "notes.txt"
    source.write_text("A Note On Something\n\n" + "Body prose. " * 40, encoding="utf-8")

    with sf() as s:
        paper_id = register_document(s, source).paper.id
        s.commit()

    run_stage(sf, JobKind.PARSE, settings, handle_parse)

    with sf() as s:
        queued = (
            s.query(Job)
            .filter(Job.paper_id == paper_id, Job.kind == JobKind.METADATA)
            .count()
        )
        assert queued == 1
