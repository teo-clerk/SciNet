"""Reading every format that is not a PDF.

The tiered PDF router exists because a PDF may or may not carry a usable text
layer, and finding out is expensive. None of that applies here. A ``.txt``,
``.md``, ``.docx`` or ``.epub`` either opens or it does not, and a ``.djvu``
either carries the OCR text its scanner saved or it carries nothing this side
of running OCR ourselves. So these bypass escalation entirely and report tier
0 — the tier field records how much work the text cost, and this is the
cheapest text there is.

Every extractor returns the same ``ParseResult`` the PDF tiers return, so the
rest of the pipeline — metadata, embedding, projection, tagging — cannot tell
where a document came from.
"""

from __future__ import annotations

import logging
from pathlib import Path

from app.core.paths import DocumentKind
from app.services.parse.tier0_pymupdf import ParseResult

logger = logging.getLogger(__name__)

PARSER_VERSION = "1"
# Roughly a printed page. Used only so page_count means something comparable to
# a PDF's; nothing downstream depends on it being exact.
CHARS_PER_PAGE = 1800


class UnreadableDocument(Exception):
    """The file cannot be turned into text at all."""


def _page_count(text: str) -> int:
    return max(1, round(len(text) / CHARS_PER_PAGE))


def read_text_file(path: Path) -> str:
    """Decode a plain text or Markdown file.

    UTF-8 first because that is what the format is meant to be. The fallbacks
    exist because a personal library accumulates files from everywhere, and
    refusing to read a paper over one bad byte helps nobody — a replacement
    character in one word is a far better outcome than a missing document.
    """
    raw = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    logger.warning(
        "%s is not valid UTF-8 or cp1252; decoding with replacement", path.name
    )
    return raw.decode("utf-8", errors="replace")


def read_docx(path: Path) -> str:
    """Flatten a .docx into Markdown.

    Word's heading styles are the only structure worth keeping: the metadata
    stage reads ``#`` headings to find the title and abstract, so dropping them
    would cost every .docx its title. Tables are rendered as rows of cells
    rather than pipe tables — good enough to embed, and not worth the fidelity.
    """
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise UnreadableDocument("python-docx is not installed") from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001 - python-docx raises many types
        raise UnreadableDocument(f"could not open {path.name}: {exc}") from exc

    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style.startswith("heading"):
            # "Heading 2" -> "##". Word's Title style is the document title.
            digits = "".join(c for c in style if c.isdigit())
            level = int(digits) if digits else 1
            lines.append(f"{'#' * min(level, 6)} {text}")
        elif style == "title":
            lines.append(f"# {text}")
        else:
            lines.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    body = "\n\n".join(lines).strip()
    if not body:
        raise UnreadableDocument(f"{path.name} contains no extractable text")
    return body


def _read_plain(path: Path) -> tuple[str, int | None]:
    return read_text_file(path), None


def _read_docx(path: Path) -> tuple[str, int | None]:
    return read_docx(path), None


def _read_epub(path: Path) -> tuple[str, int | None]:
    from app.services.parse.ebook import UnreadableBook, read_epub

    try:
        return read_epub(path)
    except UnreadableBook as exc:
        raise UnreadableDocument(str(exc)) from exc


def _read_mobi(path: Path) -> tuple[str, int | None]:
    from app.services.parse.ebook import UnreadableBook, read_mobi

    try:
        return read_mobi(path)
    except UnreadableBook as exc:
        raise UnreadableDocument(str(exc)) from exc


def _read_djvu(path: Path) -> tuple[str, int | None]:
    from app.services.parse.djvu import NotADjVu, read_djvu_text

    try:
        text, pages = read_djvu_text(path)
    except NotADjVu as exc:
        raise UnreadableDocument(str(exc)) from exc
    if not text.strip():
        # Images with no OCR layer. Reporting this as an empty document would
        # put a blank node on the map; a failure is something the operator can
        # act on, by running the file through OCR themselves.
        raise UnreadableDocument(
            f"{path.name} has no text layer — it is a scan that was never "
            "run through OCR"
        )
    return text, pages


READERS = {
    DocumentKind.TEXT: _read_plain,
    DocumentKind.MARKDOWN: _read_plain,
    DocumentKind.DOCX: _read_docx,
    DocumentKind.EPUB: _read_epub,
    DocumentKind.MOBI: _read_mobi,
    DocumentKind.DJVU: _read_djvu,
}

PARSER_NAMES = {
    DocumentKind.TEXT: "utf8-text",
    DocumentKind.MARKDOWN: "markdown-passthrough",
    DocumentKind.DOCX: "python-docx",
    DocumentKind.EPUB: "epub-xhtml",
    DocumentKind.MOBI: "mupdf-mobi",
    DocumentKind.DJVU: "djvu-textlayer",
}


def parse_text_document(path: Path | str, kind: DocumentKind) -> ParseResult:
    """Read a non-PDF document into the pipeline's common shape."""
    path = Path(path)
    reader = READERS.get(kind)
    if reader is None:
        raise UnreadableDocument(f"no reader for {kind}")

    text, pages = reader(path)
    text = text.strip()
    if not text:
        raise UnreadableDocument(f"{path.name} is empty")

    return ParseResult(
        markdown=text,
        tier=0,
        parser=PARSER_NAMES[kind],
        parser_version=PARSER_VERSION,
        # Formats with real pages report them; the rest are measured in text.
        page_count=pages if pages else _page_count(text),
    )
